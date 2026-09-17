"""
tests/test_certified_reports_and_pe_stamp.py
=============================================
Automated test suite verifying:
1. Professional Engineer (PE) Digital Regulatory Stamp and SHA-256 signatures (api/pe_stamp.py).
2. Certified IEEE Std 3002.7, IEC 60909, and IEEE 1584 PDF report generation.
3. Multi-sheet Excel workbook export with engineering standards compliance tables.
4. Export and Reports REST API endpoints (/api/v1/reports, /api/v1/exports, /api/v1/export/{id}/pdf|excel).
5. Benchmark virtual project support (ieee-9bus-wscc, ieee-14bus-feeder).
6. Production parity on Hugging Face Spaces entrypoint (hf-space/app.py).
"""

from __future__ import annotations

import io
import uuid

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from api.dependencies import CurrentUser, get_api_key, get_current_user_from_header
from api.export import exports_router, reports_router
from api.export import router as export_router
from api.pe_stamp import PEStamp, PEStampRecord, generate_pe_stamp, verify_pe_stamp
from api.rbac import require_permission
from api.services.export_generator import generate_excel_export, generate_pdf_export

# Test User
TEST_USER = CurrentUser(
    user_id="usr_lead_eng_01",
    username="ahmed_elbaz",
    email="ahmed.elbaz@engineering.com",
    role="admin",
    tenant_id="default",
)


# Mock Study Data matching IEEE 3002.7 / IEC 60909 / IEEE 1584
MOCK_STUDIES = [
    {
        "study_type": "load_flow",
        "created_at": "2026-09-17T12:00:00Z",
        "results": {
            "converged": True,
            "iterations": 4,
            "summary": {
                "total_generation_mw": 319.64,
                "total_load_mw": 315.0,
                "total_losses_mw": 4.64,
            },
            "buses": [
                {"name": "Bus 1 (Slack)", "voltage_pu": 1.040, "angle_deg": 0.0, "nominal_kv": 16.5, "type": "slack"},
                {"name": "Bus 2 (Gen 2)", "voltage_pu": 1.025, "angle_deg": 9.28, "nominal_kv": 18.0, "type": "pv"},
                {"name": "Bus 3 (Gen 3)", "voltage_pu": 1.025, "angle_deg": 4.66, "nominal_kv": 13.8, "type": "pv"},
                {"name": "Bus 4", "voltage_pu": 1.026, "angle_deg": -2.22, "nominal_kv": 230.0, "type": "pq"},
                {"name": "Bus 5", "voltage_pu": 0.996, "angle_deg": -3.99, "nominal_kv": 230.0, "type": "pq"},
                {"name": "Bus 6", "voltage_pu": 1.013, "angle_deg": -3.69, "nominal_kv": 230.0, "type": "pq"},
                {"name": "Bus 7", "voltage_pu": 1.026, "angle_deg": 3.72, "nominal_kv": 230.0, "type": "pq"},
                {"name": "Bus 8", "voltage_pu": 1.016, "angle_deg": 0.73, "nominal_kv": 230.0, "type": "pq"},
                {"name": "Bus 9", "voltage_pu": 1.032, "angle_deg": 1.97, "nominal_kv": 230.0, "type": "pq"},
            ],
            "branches": [
                {"from_bus": "Bus 1", "to_bus": "Bus 4", "loading_pct": 54.2, "p_loss_mw": 0.3},
                {"from_bus": "Bus 4", "to_bus": "Bus 5", "loading_pct": 42.1, "p_loss_mw": 1.1},
            ],
        },
    },
    {
        "study_type": "short_circuit",
        "created_at": "2026-09-17T12:05:00Z",
        "results": {
            "standard": "IEC 60909",
            "method": "complete_equivalent_source",
            "summary": {
                "max_fault_current_ka": 28.45,
                "max_fault_bus": "Bus 4 (230kV Grid)",
            },
            "faults": [
                {"bus": "Bus 1 (16.5kV)", "ik_ss": 24.2, "ip": 61.5, "sk": 691.0, "x_r": 18.5},
                {"bus": "Bus 4 (230kV)", "ik_ss": 28.45, "ip": 72.8, "sk": 11330.0, "x_r": 22.1},
                {"bus": "Bus 7 (230kV)", "ik_ss": 21.3, "ip": 54.2, "sk": 8485.0, "x_r": 19.4},
            ],
        },
    },
    {
        "study_type": "arc_flash",
        "created_at": "2026-09-17T12:10:00Z",
        "results": {
            "standard": "IEEE 1584-2018",
            "locations": [
                {"equipment": "Switchgear SWG-1", "voltage_kv": 13.8, "incident_energy_cal_cm2": 6.8, "arc_flash_boundary_mm": 1250, "ppe_category": "2"},
                {"equipment": "MCC-4A Bus", "voltage_kv": 0.48, "incident_energy_cal_cm2": 3.4, "arc_flash_boundary_mm": 680, "ppe_category": "1"},
            ],
        },
    },
]


# ===========================================================================
# 1. Professional Engineer (PE) Digital Seal & Signature Unit Tests
# ===========================================================================

def test_pe_stamp_signature_and_verification():
    """Verify that PEStamp signs data deterministically and verifies authenticity."""
    payload = {"study": "load_flow", "project": "Cairo West Substation"}
    stamp = PEStamp.sign_study(
        study_data=payload,
        engineer_name="Eng. Ahmed Elbaz, PE",
        license_id="PE-EE-2026-0915",
        jurisdiction="Egypt & North America (NCEES Model Law)",
        role="Principal Power Systems Engineer",
    )

    assert isinstance(stamp, PEStampRecord)
    assert stamp.engineer_name == "Eng. Ahmed Elbaz, PE"
    assert stamp.license_id == "PE-EE-2026-0915"
    assert len(stamp.signature_sha256) == 64  # SHA-256 hex length
    assert "IEEE Std 3002.7" in stamp.standards_statement
    assert "IEC 60909" in stamp.standards_statement

    # Verification passes with untouched payload
    assert stamp.verify(payload) is True

    # Verification fails if payload is tampered
    tampered_payload = {"study": "load_flow", "project": "Tampered Substation"}
    assert stamp.verify(tampered_payload) is False


def test_generate_pe_stamp_convenience_function():
    """Verify helper function returns dictionary with seal fields."""
    stamp_dict = generate_pe_stamp(MOCK_STUDIES[0]["results"])
    assert "engineer_name" in stamp_dict
    assert "signature_sha256" in stamp_dict
    assert "certified_date" in stamp_dict
    assert verify_pe_stamp(MOCK_STUDIES[0]["results"], stamp_dict["signature_sha256"]) is True


# ===========================================================================
# 2. Certified PDF Report Generator Unit Tests
# ===========================================================================

def test_generate_pdf_export_valid_pdf_format():
    """Verify generate_pdf_export produces a well-formed PDF with PE seal."""
    pdf_bytes = generate_pdf_export("IEEE 9-Bus WSCC System", MOCK_STUDIES)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 2000  # Detailed PDF is multi-page and significant

    # Empty studies fallback
    empty_pdf = generate_pdf_export("Empty Test Network", [])
    assert empty_pdf.startswith(b"%PDF")


# ===========================================================================
# 3. Certified Multi-Sheet Excel Workbook Unit Tests
# ===========================================================================

def test_generate_excel_export_sheets_and_content():
    """Verify generate_excel_export generates multi-sheet workbook with IEEE & IEC tables."""
    excel_bytes = generate_excel_export("IEEE 9-Bus WSCC System", MOCK_STUDIES)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    # Check required sheets
    assert "PE Stamp & Summary" in sheet_names
    assert "Load Flow (IEEE 3002.7)" in sheet_names
    assert "Short Circuit (IEC 60909)" in sheet_names

    # Check PE Stamp sheet content
    summary_ws = wb["PE Stamp & Summary"]
    found_pe = False
    for row in summary_ws.iter_rows(values_only=True):
        for cell in row:
            if cell and "Ahmed Elbaz" in str(cell):
                found_pe = True
                break
    assert found_pe, "Engineer name not found in Excel PE Stamp sheet"

    # Check Load Flow bus data
    lf_ws = wb["Load Flow (IEEE 3002.7)"]
    bus_names = [row[0] for row in lf_ws.iter_rows(min_row=2, values_only=True) if row[0]]
    assert any("Bus 1" in str(b) for b in bus_names)
    assert any("Bus 4" in str(b) for b in bus_names)

    # Check Short Circuit fault data
    sc_ws = wb["Short Circuit (IEC 60909)"]
    fault_buses = [row[0] for row in sc_ws.iter_rows(min_row=2, values_only=True) if row[0]]
    assert any("Bus 4" in str(b) for b in fault_buses)


# ===========================================================================
# 4. API Endpoints Integration Tests (TestClient)
# ===========================================================================

@pytest.fixture
def api_client():
    """Test FastAPI application with export routers and auth overrides."""
    app = FastAPI(title="Test AhmedETAP API")
    app.include_router(export_router)
    app.include_router(reports_router)
    app.include_router(exports_router)

    def _auth_override():
        return TEST_USER

    app.dependency_overrides[get_current_user_from_header] = _auth_override
    app.dependency_overrides[get_api_key] = _auth_override
    app.dependency_overrides[require_permission("export", "create")] = _auth_override
    app.dependency_overrides[require_permission("export", "list")] = _auth_override
    app.dependency_overrides[require_permission("export", "read")] = _auth_override

    return TestClient(app)


def test_get_reports_list(api_client):
    """GET /api/v1/reports returns report listing including canonical IEEE benchmarks."""
    res = api_client.get("/api/v1/reports")
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    assert len(data) >= 2

    # Check benchmark presence
    names = [r["name"] for r in data]
    assert any("IEEE 9-Bus" in n for n in names)
    assert any("IEEE 14-Bus" in n for n in names)

    # Check structure
    sample = data[0]
    assert "id" in sample
    assert "format" in sample
    assert "download_url" in sample
    assert sample["download_url"].startswith("/api/v1/export/")


def test_get_exports_list(api_client):
    """GET /api/v1/exports returns export history listing."""
    res = api_client.get("/api/v1/exports")
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    assert len(data) >= 1

    sample = data[0]
    assert "name" in sample
    assert "size" in sample
    assert "download_url" in sample


def test_export_benchmark_pdf_stream(api_client):
    """GET /api/v1/export/ieee-9bus-wscc/pdf generates certified PDF on the fly."""
    res = api_client.get("/api/v1/export/ieee-9bus-wscc/pdf")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
    assert len(res.content) > 1000


def test_export_benchmark_excel_stream(api_client):
    """GET /api/v1/export/ieee-9bus-wscc/excel generates certified Excel on the fly."""
    res = api_client.get("/api/v1/export/ieee-9bus-wscc/excel")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "PE Stamp & Summary" in wb.sheetnames
    assert "Load Flow (IEEE 3002.7)" in wb.sheetnames


def test_export_benchmark_post_pdf(api_client):
    """POST /api/v1/export/ieee-14bus-feeder/pdf returns certified PDF stream."""
    res = api_client.post("/api/v1/export/ieee-14bus-feeder/pdf")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")


# ===========================================================================
# 5. Production Parity Test on Hugging Face Spaces Entry Point
# ===========================================================================

def test_hf_space_router_parity():
    """Verify hf-space/app.py registers all core production routers and studies/re-run."""
    # Import hf-space app dynamically
    import importlib.util
    import os

    app_path = os.path.abspath("hf-space/app.py")
    assert os.path.exists(app_path), "hf-space/app.py must exist"

    spec = importlib.util.spec_from_file_location("hf_space_app", app_path)
    hf_mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(hf_mod)

    hf_app = hf_mod.app
    assert hf_app is not None

    route_paths = {r.path for r in hf_app.routes if hasattr(r, "path")}

    # 1. Parity routers
    assert "/api/v1/reports" in route_paths
    assert "/api/v1/exports" in route_paths
    assert "/api/v1/export/{project_id}/pdf" in route_paths
    assert "/api/v1/export/{project_id}/excel" in route_paths
    assert "/api/v1/export/formats" in route_paths
    assert "/api/v1/approvals" in route_paths
    assert "/api/v1/feature-flags" in route_paths
    assert "/api/v1/templates" in route_paths

    # 2. Studies re-run endpoint
    assert "/api/v1/studies/re-run" in route_paths
